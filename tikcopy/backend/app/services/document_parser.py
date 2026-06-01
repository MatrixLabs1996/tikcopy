"""Extrai texto puro de arquivos PDF, DOCX e TXT enviados pelo usuário."""

from pathlib import Path


def extract_text(file_path: str) -> str:
    """Detecta o tipo do arquivo pela extensão e extrai o texto."""
    ext = Path(file_path).suffix.lower()
    if ext == ".pdf":
        return _extract_pdf(file_path)
    if ext == ".docx":
        return _extract_docx(file_path)
    if ext in (".txt", ".md"):
        return Path(file_path).read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"Tipo de arquivo não suportado: {ext}. Use PDF, DOCX, TXT ou MD.")


def _extract_pdf(file_path: str) -> str:
    from pypdf import PdfReader
    reader = PdfReader(file_path)
    parts = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            parts.append(text)
    return "\n\n".join(parts).strip()


def _extract_docx(file_path: str) -> str:
    from docx import Document
    doc = Document(file_path)
    lines = []
    for para in doc.paragraphs:
        if para.text.strip():
            lines.append(para.text)
    # Inclui texto de tabelas também
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    lines.append(cell.text)
    return "\n".join(lines).strip()
